import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import glob
import os
import re
import sys

if sys.stdout is not None and hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def is_scratched_row(page, y_center):
    """
    Detects scratched / strikethrough / crossed-out rows with high performance.
    Only checks horizontal vector lines/edges passing through the row's text center band.
    """
    lines_to_check = page.lines + page.edges
    for obj in lines_to_check:
        y0 = min(obj['top'], obj['bottom'])
        y1 = max(obj['top'], obj['bottom'])
        y_mid = (y0 + y1) / 2.0
        dist = abs(y_mid - y_center)
        
        if dist <= 5.5:
            x0 = min(obj['x0'], obj['x1'])
            x1 = max(obj['x0'], obj['x1'])
            dx = x1 - x0
            if (x0 <= 180 and x1 >= 240) or (dx > 70 and x0 < 300 and x1 > 150):
                return True
    return False

def parse_all_pdfs_with_dimension(pdf_files, progress_callback=None):
    """
    Parsing module specifically for PDFs WITH dimension table headers (A, B, C, D, etc.)
    Extracts table contents directly by matching text against header column boundaries,
    preserving exact 1-to-1 row layout matching the PDF (including section header rows).
    """
    all_extracted_rows = []
    
    total_pages = 0
    for pdf_path in pdf_files:
        try:
            with pdfplumber.open(pdf_path) as p_doc:
                total_pages += len(p_doc.pages)
        except Exception:
            pass
    if total_pages == 0:
        total_pages = 1
        
    processed_pages = 0
    
    for pdf_path in pdf_files:
        pdf_name = os.path.basename(pdf_path)
        print(f"Processing {pdf_name} (with dimension headers)...")
        
        with pdfplumber.open(pdf_path) as pdf:
            found_bbs_table = False
            for p_idx, page in enumerate(pdf.pages):
                processed_pages += 1
                if progress_callback:
                    try:
                        progress_callback(processed_pages, total_pages, f"Processing {pdf_name} (Page {p_idx+1}/{len(pdf.pages)})")
                    except Exception:
                        pass
                        
                page_text = page.extract_text() or ""
                has_bbs = "BAR BENDING SCHEDULE" in page_text.upper()
                
                # Early stop: if we already processed BBS pages and hit a CAD drawing/image without BBS header, stop scanning this PDF!
                if not has_bbs:
                    if found_bbs_table:
                        print(f"Non-BBS drawing/image found on page {p_idx+1}. Stopping scan for {pdf_name}!")
                        break
                    else:
                        continue
                        
                found_bbs_table = True
                    
                words = page.extract_words()
                
                # Standard default column boundaries (x0 coordinates in pt)
                col_bounds = {
                    'member': (0, 90),
                    'mark': (90, 135),
                    'type': (135, 147),
                    'size': (147, 170),
                    'mbrs': (170, 195),
                    'each': (195, 230),
                    'tot': (230, 260),
                    'code': (260, 290),
                    'a': (380, 415),
                    'b': (415, 445),
                    'c': (445, 480),
                    'd': (480, 510),
                    'length': (510, 540),
                    'weight': (540, 600),
                }
                
                # Dynamically adjust column boundaries if page headers are present
                header_words = [w for w in words if 170 <= w['top'] <= 225]
                a_word = next((w for w in header_words if w['text'].strip() == 'A'), None)
                b_word = next((w for w in header_words if w['text'].strip() == 'B'), None)
                c_word = next((w for w in header_words if w['text'].strip() == 'C'), None)
                d_word = next((w for w in header_words if w['text'].strip() == 'D'), None)
                tot_len_word = next((w for w in header_words if 'LENGTH' in w['text'].upper() or '总长' in w['text']), None)
                wt_word = next((w for w in header_words if 'WEIGHT' in w['text'].upper() or '重量' in w['text']), None)
                
                if a_word and b_word and c_word and d_word:
                    col_bounds['a'] = (a_word['x0'] - 15, (a_word['x0'] + b_word['x0'])/2)
                    col_bounds['b'] = ((a_word['x0'] + b_word['x0'])/2, (b_word['x0'] + c_word['x0'])/2)
                    col_bounds['c'] = ((b_word['x0'] + c_word['x0'])/2, (c_word['x0'] + d_word['x0'])/2)
                    if tot_len_word:
                        col_bounds['d'] = ((c_word['x0'] + d_word['x0'])/2, (d_word['x0'] + tot_len_word['x0'])/2)
                        if wt_word:
                            col_bounds['length'] = ((d_word['x0'] + tot_len_word['x0'])/2, (tot_len_word['x0'] + wt_word['x0'])/2)
                            col_bounds['weight'] = ((tot_len_word['x0'] + wt_word['x0'])/2, 600)
                
                body_words = [w for w in words if 215 < w['top'] < 785]
                if not body_words:
                    continue
                    
                body_words.sort(key=lambda w: w['top'])
                
                lines = []
                curr_line = []
                curr_top = None
                for w in body_words:
                    if curr_top is None or abs(w['top'] - curr_top) < 9:
                        curr_line.append(w)
                        if curr_top is None:
                            curr_top = w['top']
                    else:
                        lines.append(curr_line)
                        curr_line = [w]
                        curr_top = w['top']
                if curr_line:
                    lines.append(curr_line)
                    
                for line in lines:
                    line.sort(key=lambda w: w['x0'])
                    full_str = " ".join([w['text'] for w in line]).strip()
                    
                    # Filter out headers & footers
                    if any(k in full_str.upper() for k in ["REBAR SIZE", "BAR BENDING", "TOTAL WEIGHT", "ELEMENT MARKING"]):
                        continue
                    if "MEMBER" in full_str.upper() and "MARK" in full_str.upper():
                        continue
                    if any(k in full_str.upper() for k in ["STOREY", "BLOCK", "PROJECT"]):
                        continue
                        
                    # Filter scratched / crossed-out rows
                    table_words = [w for w in line if w['x0'] < 310]
                    y_center = min(w['top'] for w in table_words) if table_words else min(w['top'] for w in line)
                    if is_scratched_row(page, y_center):
                        continue
                        
                    def get_col_text(bounds):
                        w_list = [w['text'] for w in line if bounds[0] <= w['x0'] < bounds[1]]
                        return " ".join(w_list).strip()
                        
                    mbr_str = get_col_text(col_bounds['member'])
                    mark_str = get_col_text(col_bounds['mark'])
                    type_str = get_col_text(col_bounds['type'])
                    size_str = get_col_text(col_bounds['size'])
                    mbrs_str = get_col_text(col_bounds['mbrs'])
                    each_str = get_col_text(col_bounds['each'])
                    tot_str = get_col_text(col_bounds['tot'])
                    code_str = get_col_text(col_bounds['code'])
                    a_str = get_col_text(col_bounds['a'])
                    b_str = get_col_text(col_bounds['b'])
                    c_str = get_col_text(col_bounds['c'])
                    d_str = get_col_text(col_bounds['d'])
                    
                    # Skip completely empty rows
                    if not any([mbr_str, mark_str, type_str, size_str, mbrs_str, each_str, tot_str, code_str, a_str, b_str, c_str, d_str]):
                        continue
                        
                    # Split combined Type and Size if needed (e.g. H 16)
                    if not size_str and type_str:
                        m = re.search(r'([A-Za-z]+)\s*(\d+)', type_str)
                        if m:
                            type_str = m.group(1)
                            size_str = m.group(2)
                    elif not type_str and size_str:
                        m = re.search(r'([A-Za-z]+)\s*(\d+)', size_str)
                        if m:
                            type_str = m.group(1)
                            size_str = m.group(2)
                            
                    all_extracted_rows.append({
                        'Member': mbr_str,
                        'Bar Mark': mark_str,
                        'Type': type_str,
                        'Size': size_str,
                        'No. of MBRS': mbrs_str,
                        'No. of EACH': each_str,
                        'TOTAL': tot_str,
                        'SHAPE CODE': code_str,
                        'A': a_str,
                        'B': b_str,
                        'C': c_str,
                        'D': d_str,
                        'E': "",
                    })
                    
    return all_extracted_rows


def export_to_excel(all_rows, output_path):
    cols = ['Member', 'Bar Mark', 'Type', 'Size', 'No. of MBRS', 'No. of EACH', 'TOTAL', 'SHAPE CODE', 'A', 'B', 'C', 'D', 'E']
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BAR BENDING SCHEDULE"
    ws.views.sheetView[0].showGridLines = True
    
    font_header = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    font_data = Font(name='Segoe UI', size=10)
    
    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    
    thin_side = Side(style='thin', color='D9D9D9')
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    ws.append(cols)
    for col_idx in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
    ws.row_dimensions[1].height = 26
    
    for row_idx, r in enumerate(all_rows, start=2):
        row_vals = [r.get(c, "") for c in cols]
        ws.append(row_vals)
        ws.row_dimensions[row_idx].height = 20
        
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_data
            cell.border = border_cell
            col_name = cols[col_idx - 1]
            if col_name in ['Member', 'Bar Mark']:
                cell.alignment = align_left
            else:
                cell.alignment = align_center
                
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    wb.save(output_path)
    print(f"\nSUCCESS: Exported {len(all_rows)} rows to '{output_path}'")

if __name__ == "__main__":
    pdf_files = sorted(glob.glob("*.pdf"))
    print(f"Found {len(pdf_files)} PDF files: {pdf_files}")
    dim_files = [f for f in pdf_files if "H389" in f] # test run on dimension PDF
    target_files = dim_files if dim_files else pdf_files
    data = parse_all_pdfs_with_dimension(target_files)
    export_to_excel(data, "BAR_BENDING_SCHEDULE_WITH_DIMENSION_EXTRACTED.xlsx")
